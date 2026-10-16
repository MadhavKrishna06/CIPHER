"""Turn course documents into (text, metadata) pages.

One record per page/slide so citations can point at "Lecture 2, slide 14".
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path


@dataclass
class Page:
    text: str
    source: str   # file name shown in citations
    page: int     # 1-based page or slide number as the PDF counts it
    printed: int | None = None   # the number printed on the page, when it differs


# Sub-folders of data/raw that say what a document is. A citation must read
# "slide 9" for a lecture deck and "page 447" for a textbook, and a PDF extension
# alone cannot tell them apart.
_SLIDE_DIRS = {"slides", "lectures", "decks"}
_PAGE_DIRS = {"textbooks", "books", "papers", "rfcs"}


def doc_kind(rel_path: str, name: str) -> str:
    """Return "slide" or "page" for citations.

    The sub-folder under data/raw decides when there is one. Flat layouts fall back
    to the file name, so an existing knowledge base keeps citing correctly.
    """
    top = rel_path.split("/", 1)[0].lower() if "/" in rel_path else ""
    if top in _SLIDE_DIRS:
        return "slide"
    if top in _PAGE_DIRS:
        return "page"
    low = name.lower()
    if low.endswith(".pptx") or "slide" in low or "lecture" in low:
        return "slide"
    return "page"


# PDF extraction keeps the hyphen a typesetter used to break a word across lines, so
# "different" arrives as "differ- ent" and embeds as two meaningless tokens.
_LINE_BREAK_HYPHEN = re.compile(r"(\w)-\s+(\w)")


def _clean(text: str) -> str:
    return _LINE_BREAK_HYPHEN.sub(r"\1\2", " ".join(text.split()))


# --- running headers, footers and printed page numbers -----------------------
# A textbook prints a header on every page: "118 CHAPTER 4 / KEY DISTRIBUTION AND
# USER AUTHENTICATION" on one side, "4.3 /KERBEROS 119" on the other. Extraction
# glues it to the body, so chunks began "4.3 /KERBEROS 119 is 28 x 5 = 1280 minutes"
# and the header text competed with real content during retrieval.
#
# The header also carries the number the reader actually sees. A PDF page index is
# not that number, because front matter is unnumbered: Stallings page 136 of the file
# prints 118, and Kurose page 219 prints 191. Citing the file index sends a reader
# to the wrong page, which matters because citations are the graded output.

# A word that can belong to a running header: all caps, a section number, or a
# separator glyph. Anything containing lower case ends the header.
_SEP = r"[/|•–—∙�.\-]"
# Extraction glues the separator to the next word, so "/KERBEROS" is one token.
_HEADER_WORD = re.compile(
    rf"^(?:{_SEP}*[A-Z][A-Z0-9'&.,:\-]*|{_SEP}*\d+(?:\.\d+)*|{_SEP}+)$"
)

# Typesetting footer left in the Kurose PDF by InDesign. Holds the printed number.
_INDD_FOOTER = re.compile(
    r"\bM\d{2}_[A-Z0-9_]+\.indd\s+(\d{1,4})\s+\d{1,2}/\d{1,2}/\d{2,4}\s+"
    r"\d{1,2}:\d{2}\s*[AaPp]\.?[Mm]\.?\s*$"
)

# A header run is short. Anything longer is body text that happens to be shouted.
_MAX_HEADER_WORDS = 14


def _header_span(words: list[str]) -> int:
    """How many leading words form a running header. 0 means none."""
    n = 0
    while n < min(len(words), _MAX_HEADER_WORDS) and _HEADER_WORD.match(words[n]):
        n += 1
    # A single capitalised word is just a sentence opening, not a header.
    return n if n >= 2 else 0


def _printed_from_header(words: list[str], span: int) -> int | None:
    """The page number inside a header run, which sits at one end or the other."""
    if span == 0:
        return None
    for candidate in (words[0], words[span - 1]):
        if candidate.isdigit() and 1 <= int(candidate) <= 2000:
            return int(candidate)
    return None


def _page_signals(text: str) -> tuple[int | None, int, str]:
    """Return the printed page number, the header word count, and the footer-free text."""
    footer = _INDD_FOOTER.search(text)
    printed_footer = int(footer.group(1)) if footer else None
    body = text[: footer.start()].rstrip() if footer else text

    words = body.split()
    span = _header_span(words)
    printed_header = _printed_from_header(words, span)

    return printed_footer or printed_header, span, body


def _detect_offset(raw_pages: list[str]) -> int | None:
    """The constant difference between PDF page index and printed page number.

    Taking the mode rather than the first match tolerates pages where extraction
    garbled the header, chapter openings that print no number, and figure pages.
    """
    offsets: Counter[int] = Counter()
    for i, text in enumerate(raw_pages, start=1):
        printed, _, _ = _page_signals(text)
        if printed is not None and 0 < i - printed < 200:
            offsets[i - printed] += 1

    if not offsets:
        return None
    offset, hits = offsets.most_common(1)[0]
    # Require real agreement. A handful of coincidental matches is not an offset.
    return offset if hits >= max(5, len(raw_pages) // 10) else None


def _strip_header(text: str) -> str:
    """Remove the running header and the typesetting footer from one page."""
    _, span, body = _page_signals(text)
    if span:
        body = " ".join(body.split()[span:])
    return body.strip()


def load_pdf(path: Path) -> list[Page]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    raw = [_clean(p.extract_text() or "") for p in reader.pages]
    offset = _detect_offset(raw)

    pages = []
    for i, text in enumerate(raw, start=1):
        body = _strip_header(text)
        if not body:
            continue
        printed = i - offset if offset is not None and i - offset >= 1 else None
        pages.append(Page(body, path.name, i, printed))
    return pages


def load_pptx(path: Path) -> list[Page]:
    from pptx import Presentation

    prs = Presentation(str(path))
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
        text = _clean("\n".join(parts))
        if text:
            pages.append(Page(text, path.name, i))
    return pages


def load_docx(path: Path) -> list[Page]:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    text = _clean("\n".join(parts))
    return [Page(text, path.name, 1)] if text else []


def load_text(path: Path) -> list[Page]:
    text = _clean(path.read_text(encoding="utf-8", errors="ignore"))
    return [Page(text, path.name, 1)] if text else []


_LOADERS = {
    ".pdf": load_pdf,
    ".pptx": load_pptx,
    ".docx": load_docx,
    ".txt": load_text,
    ".md": load_text,
}


def load_document(path: Path) -> list[Page]:
    loader = _LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"unsupported file type: {path.suffix}")
    return loader(path)
